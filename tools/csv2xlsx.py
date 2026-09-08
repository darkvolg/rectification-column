#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Журнал ректификации: CSV → Excel.

Делает из скачанного журнала нормальный .xlsx: шапка «сырьё и условия»,
блок расчёта фракций ЖИВЫМИ ФОРМУЛАМИ (поменял крепость — пересчиталось
всё), таблица с заливкой по фазам, закреплённая шапка, фильтр.

Вторым листом можно подшить подробный лог контроллера (log.csv) —
тот, что пишется раз в 20 секунд.

Запуск:
    python csv2xlsx.py Журнал_ректификации_2026-09-01.csv
    python csv2xlsx.py журнал.csv --log log.csv -o Погон_01.xlsx

Нужен openpyxl:
    pip install openpyxl
"""

__version__ = "1.0.0"

import argparse
import csv
import io
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Нужен openpyxl:  pip install openpyxl")

# Цвета фаз — те же, что подсвечивают строку на странице журнала
PHASE_FILL = {
    "Прогрев":      "E8EAED",
    "Стабилизация": "DCE9F2",
    "Головы":       "F8DDD9",
    "Подголовники": "FBEED5",
    "Тело":         "DEEDE4",
    "Хвосты":       "F0E2D2",
    "Стоп":         "D9D9D9",
}

HDR_FILL = "1F3864"      # тёмно-синий, как в бумажном журнале
CALC_FILL = "EDF3E7"     # расчётные ячейки — светло-зелёные
IN_FILL = "FFF6D9"       # то, что вводится руками — жёлтые

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_csv(path):
    """Читает CSV страницы: BOM, разделитель «;», десятичная запятая."""
    with io.open(path, encoding="utf-8-sig", newline="") as fh:
        return [r for r in csv.reader(fh, delimiter=";")]


def num(s):
    """'78,256' → 78.256. Не число — возвращаем как есть."""
    if s is None:
        return None
    t = str(s).strip()
    if not t:
        return None
    if re.fullmatch(r"[-+]?\d+([.,]\d+)?", t):
        v = float(t.replace(",", "."))
        return int(v) if v == int(v) else v
    return t


def find_head(rows):
    """Достаёт значения шапки из выгрузки страницы."""
    want = {
        "Объём СС": "vol", "Крепость СС": "abv", "Давление": "press",
        "t° начала в кубе": "tstart", "% голов от АС": "head",
        "% хвостов от АС": "tail",
    }
    out = {}
    for r in rows:
        if len(r) >= 2 and r[0].strip() in want:
            out[want[r[0].strip()]] = num(r[1])
    return out


def find_table(rows):
    """Находит строку заголовков таблицы и всё, что после неё."""
    for i, r in enumerate(rows):
        if r and r[0].strip() == "№":
            return r, [x for x in rows[i + 1:] if any(c.strip() for c in x)]
    return None, []


def style_title(ws, cell, text, span):
    ws[cell] = text
    ws[cell].font = Font(bold=True, size=11, color="FFFFFF")
    ws[cell].fill = PatternFill("solid", fgColor=HDR_FILL)
    ws[cell].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(span)


def build(journal, logfile, out):
    rows = read_csv(journal)
    head = find_head(rows)
    thead, tbody = find_table(rows)
    if not thead:
        sys.exit("В файле не нашлась таблица журнала — это точно выгрузка со страницы?")

    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал"

    # ---------------- заголовок ----------------
    style_title(ws, "B1", "ЖУРНАЛ РЕКТИФИКАЦИИ", "B1:M1")
    ws.row_dimensions[1].height = 24

    # ---------------- сырьё и условия ----------------
    style_title(ws, "B3", "СЫРЬЁ И УСЛОВИЯ", "B3:D3")
    left = [
        ("Объём СС:",          head.get("vol", 22000), "мл",         None),
        ("Крепость СС:",       head.get("abv", 30),    "%",          None),
        ("Давление:",          head.get("press", 760), "мм рт.ст.",  None),
        # Формула живая: сменилось давление — поехала точка кипения,
        # а за ней все ΔT в таблице. Ровно так же считает прибор.
        ("t° кипения спирта:", "=78.15+(C6-760)*0.037", "°C",        "calc"),
        ("t° начала в кубе:",  head.get("tstart", ""), "°C",         None),
        ("% голов от АС:",     head.get("head", 15),   "%",          None),
        ("% хвостов от АС:",   head.get("tail", 5),    "%",          None),
    ]
    for i, (lab, val, unit, kind) in enumerate(left):
        r = 4 + i
        ws.cell(r, 2, lab).font = Font(bold=True, size=10)
        ws.cell(r, 2).alignment = Alignment(horizontal="right")
        c = ws.cell(r, 3, val)
        c.font = Font(bold=True, size=11)
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor=CALC_FILL if kind == "calc" else IN_FILL)
        c.border = BOX
        if lab.startswith("t°"):
            c.number_format = "0.000" if kind == "calc" else "0.0"
        ws.cell(r, 4, unit).font = Font(size=9, color="808080")

    # ---------------- расчёт фракций ----------------
    style_title(ws, "F3", "РАСЧЁТ ФРАКЦИЙ (формулы от АС)", "F3:H3")
    # Ссылки собраны под порядок строк слева:
    #   C4 объём, C5 крепость, C6 давление, C7 t кипения,
    #   C8 t начала, C9 % голов, C10 % хвостов
    # и справа: G4 АС, G5 головы, G6 тело, G7 хвосты.
    right = [
        ("Абсолютный спирт (АС):", "=C4*C5/100",        "мл", "calc"),
        ("Головы:",                "=G4*C9/100",        "мл", "calc"),
        ("Тело:",                  "=G4-G5-G7",         "мл", "calc"),
        ("Хвосты:",                "=G4*C10/100",       "мл", "calc"),
        ("Головохвосты итого:",    "=G5+G7",            "мл", "calc"),
        ("Доля тела от АС:",       "=IF(G4>0,G6/G4,0)", "%",  "calc"),
        ("Мощность ТЭН:",          "",                  "Вт", "in"),
    ]

    for i, (lab, val, unit, kind) in enumerate(right):
        r = 4 + i
        ws.cell(r, 6, lab).font = Font(bold=True, size=10)
        ws.cell(r, 6).alignment = Alignment(horizontal="right")
        c = ws.cell(r, 7, val)
        c.font = Font(bold=True, size=11,
                      color="1F6B3B" if kind == "calc" else "000000")
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor=CALC_FILL if kind == "calc" else IN_FILL)
        c.border = BOX
        c.number_format = "0.0%" if unit == "%" else "# ##0"
        ws.cell(r, 8, unit).font = Font(size=9, color="808080")

    # ---------------- таблица ----------------
    HR = 12
    for j, name in enumerate(thead):
        c = ws.cell(HR, 2 + j, name)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HDR_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[HR].height = 32

    for i, row in enumerate(tbody):
        r = HR + 1 + i
        phase = row[2].strip() if len(row) > 2 else ""
        for j, val in enumerate(row):
            c = ws.cell(r, 2 + j, num(val))
            c.alignment = Alignment(horizontal="left" if j == 11 else "center")
            c.border = BOX
            c.font = Font(size=10)
            if j in (3, 4, 5):
                c.number_format = "0.00"
            elif j == 10:
                c.number_format = "+0.000;-0.000;0.000"
        # Накопленное и ΔT — формулами, чтобы правка отбора пересчитала итог
        ws.cell(r, 10, "=SUM($I$%d:I%d)" % (HR + 1, r)).number_format = "# ##0"
        ws.cell(r, 12, "=IF(E%d=\"\",\"\",E%d-$C$7)" % (r, r)).number_format = \
            "+0.000;-0.000;0.000"
        if phase in PHASE_FILL:
            fill = PatternFill("solid", fgColor=PHASE_FILL[phase])
            for j in range(2, 2 + len(thead)):
                ws.cell(r, j).fill = fill

    widths = [5, 6, 9, 14, 12, 12, 11, 9, 9, 10, 11, 11, 34]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(2 + i)].width = w
    # Подписи шапки выровнены вправо и переполняются в колонку A.
    # Она пустая, но узкая — тогда длинные подписи обрезаются о край листа.
    ws.column_dimensions["A"].width = 16

    ws.freeze_panes = "B%d" % (HR + 1)
    if tbody:
        ws.auto_filter.ref = "B%d:%s%d" % (
            HR, get_column_letter(1 + len(thead)), HR + len(tbody))

    # ---------------- лог контроллера ----------------
    if logfile:
        add_log(wb, logfile)

    wb.save(out)
    return len(tbody), out


def add_log(wb, path):
    """Подробный лог с контроллера — вторым листом, как есть."""
    rows = read_csv(path)
    if not rows:
        return
    ws = wb.create_sheet("Секунды")
    for j, name in enumerate(rows[0]):
        c = ws.cell(1, 1 + j, name)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HDR_FILL)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows[1:]):
        for j, val in enumerate(row):
            ws.cell(2 + i, 1 + j, num(val))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 19
    for j in range(2, len(rows[0]) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    if len(rows) > 1:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(rows[0])), len(rows))


def main():
    ap = argparse.ArgumentParser(description="Журнал ректификации: CSV → Excel")
    ap.add_argument("journal", help="файл, скачанный со страницы «Журнал»")
    ap.add_argument("--log", help="log.csv с контроллера — вторым листом")
    ap.add_argument("-o", "--out", help="имя .xlsx (по умолчанию рядом с журналом)")
    a = ap.parse_args()

    if not os.path.exists(a.journal):
        sys.exit("Нет файла: %s" % a.journal)
    out = a.out or os.path.splitext(a.journal)[0] + ".xlsx"
    n, path = build(a.journal, a.log, out)
    print("Готово: %s, строк журнала %d" % (path, n))


if __name__ == "__main__":
    main()
